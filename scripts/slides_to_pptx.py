#!/usr/bin/env python3
"""Export the HTML deck to PowerPoint, keeping the Magic-Move transitions as PowerPoint *Morph*.

    python3 scripts/slides_to_pptx.py [slides/index.html] [slides/defense-slides.pptx]
                                      [--theme=light|dark] [--mode=native|image] [--no-morph]

--mode=native (default)  Every piece of the slide is its own PowerPoint object: each headline, kicker, label,
    bullet, table cell and chart label is an editable text box (Arial / Courier New, same size, colour, weight,
    spacing and alignment as the web deck); cards, chips, rules and table lines are shapes; the Texas Tech
    Double T is a picture; each diagram/chart (svg.fig) is a picture with its text labels lifted out as text
    boxes. The web deck's morphing elements keep their data-morph id as the shape name "!!<id>", which is how
    PowerPoint Morph matches objects between slides, so labels, cards, rail chips, bars and chart points move
    exactly where the web deck morphs them. Text may reflow by a word here and there (PowerPoint is not a
    browser); the web deck remains the reference.
--mode=image  One pixel-perfect 1920x1080 render per slide plus cut-out pictures for the morphing elements
    (not editable, but exact).

Slide N gets a 650 ms Morph transition when it shares a data-morph-group token with slide N-1, otherwise a Fade;
PowerPoint versions without Morph fall back to Fade. <aside class="notes"> becomes the speaker notes.
Requires playwright (with Chromium) and python-pptx. Re-run after editing index.html."""
import re, sys, html, json, pathlib, tempfile
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches, Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree

args = [a for a in sys.argv[1:] if not a.startswith("--")]
opt = lambda k, d: next((a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith(f"--{k}=")), d)
theme, mode, MORPH = opt("theme", "light"), opt("mode", "native"), "--no-morph" not in sys.argv
src = pathlib.Path(args[0] if len(args) > 0 else "slides/index.html").resolve()
out = pathlib.Path(args[1] if len(args) > 1 else "slides/defense-slides.pptx").resolve()
W, H = 1920, 1080
PX = 914400 * 13.333 / W                  # EMU per CSS px (13.333 in wide slide -> 1 px = 1/144 in = 0.5 pt)
PT = 0.5                                  # points per CSS px
DUR_MS = 650
SANS, MONO = "Arial", "Courier New"        # metric twins of the fonts Chromium used for the layout measurements

# ---- slide metadata straight from the HTML -----------------------------------------------------
text = src.read_text(encoding="utf-8")
body = re.search(r"<body[^>]*>(.*?)</body>", text, re.S).group(1)
sections = re.findall(r'<section class="slide[^"]*"[^>]*>.*?</section>', body, re.S)
doc_title = html.unescape(re.search(r"<title>(.*?)</title>", text, re.S).group(1).strip())
strip = lambda s: html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", s))).strip()
meta = []
for s in sections:
    head = re.match(r"<section[^>]*>", s).group(0)
    m = re.search(r'data-title="([^"]*)"', head); g = re.search(r'data-morph-group="([^"]*)"', head)
    h = re.search(r"<h1[^>]*>(.*?)</h1>", s, re.S); n = re.search(r'<aside class="notes">(.*?)</aside>', s, re.S)
    meta.append({"title": html.unescape(m.group(1)) if m else (strip(h.group(1)) if h else ""),
                 "groups": set(g.group(1).split()) if g else set(), "notes": strip(n.group(1)) if n else ""})
N = len(meta)
share = [False] + [bool(meta[i - 1]["groups"] & meta[i]["groups"]) for i in range(1, N)]

# ---- in-page extraction (native mode): walk the active slide in paint order ----------------------
EXTRACT_JS = r"""
(cfg) => {
  const slide = document.querySelector('.slide.active');
  const want = new Set(cfg.want);            // morph ids matched on a neighbouring slide (cut-outs inside figures)
  const items = []; let k = 0;
  const flag = el => { const id = 'x' + (k++); el.setAttribute('data-x', id); return id; };
  const num = v => parseFloat(v) || 0;
  const rgba = c => { const m = (c || '').match(/rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)/);
    if (!m) return null; const a = m[4] === undefined ? 1 : parseFloat(m[4]); if (a <= 0) return null;
    return { hex: [m[1], m[2], m[3]].map(v => Math.round(+v).toString(16).padStart(2, '0')).join('').toUpperCase(), a }; };
  const box = el => { const r = el.getBoundingClientRect(); return { x: r.left, y: r.top, w: r.width, h: r.height }; };
  const style = (cs, el) => ({
    size: num(cs.fontSize), color: rgba(cs.color), bold: parseInt(cs.fontWeight) >= 600, italic: cs.fontStyle === 'italic',
    mono: /mono|Menlo|Consolas|Courier/i.test(cs.fontFamily), caps: cs.textTransform === 'uppercase',
    ls: cs.letterSpacing === 'normal' ? 0 : num(cs.letterSpacing),
    grad: !!(el && el.classList && el.classList.contains('grad')),
    sup: !!(el && (el.tagName === 'SUP' || cs.verticalAlign === 'super')), sub: !!(el && el.tagName === 'SUB') });
  const INLINE = new Set(['B', 'STRONG', 'I', 'EM', 'SPAN', 'SUP', 'SUB', 'CODE', 'SMALL', 'A', 'BR']);
  const isInline = el => el.tagName === 'BR' || (INLINE.has(el.tagName) && getComputedStyle(el).display.startsWith('inline'));
  const visible = el => { const cs = getComputedStyle(el); return cs.display !== 'none' && cs.visibility !== 'hidden' && num(cs.opacity) > 0; };
  const isTextContainer = el => {
    if (!el.textContent.trim()) return false;
    for (const n of el.childNodes) if (n.nodeType === 1 && !isInline(n)) return false;
    return true; };
  function collectRuns(el, base, pre, out) {
    for (const n of el.childNodes) {
      if (n.nodeType === 3) { let t = n.textContent; if (!pre) t = t.replace(/[ \t\r\n]+/g, ' '); if (t) out.push(Object.assign({ t }, base)); }
      else if (n.nodeType === 1) {
        if (n.tagName === 'BR') { out.push({ br: true }); continue; }
        if (!visible(n)) continue;
        const cs = getComputedStyle(n), st = style(cs, n);
        if (!st.color) { st.color = base.color; st.grad = st.grad || base.grad; }   // transparent (gradient-clipped) text inherits the parent's colour
        collectRuns(n, st, pre, out);
      } } }
  function textItem(el) {
    const cs = getComputedStyle(el), b = box(el);
    const pad = [cs.paddingLeft, cs.paddingTop, cs.paddingRight, cs.paddingBottom].map(num);
    const bw = [cs.borderLeftWidth, cs.borderTopWidth, cs.borderRightWidth, cs.borderBottomWidth].map(num);
    const pre = cs.whiteSpace.startsWith('pre');
    const runs = []; collectRuns(el, style(cs, el), pre, runs);      // the container's own .grad / caps apply to its bare text
    const lh = cs.lineHeight === 'normal' ? num(cs.fontSize) * 1.2 : num(cs.lineHeight);
    const align = { start: 'left', left: 'left', center: 'center', right: 'right', end: 'right', justify: 'left' }[cs.textAlign] || 'left';
    return { type: 'text', id: flag(el), x: b.x + bw[0] + pad[0], y: b.y + bw[1] + pad[1], w: b.w - bw[0] - bw[2] - pad[0] - pad[2],
             h: b.h - bw[1] - bw[3] - pad[1] - pad[3], align, lh, pre, wrap: cs.whiteSpace !== 'nowrap' && !pre, runs,
             morph: el.dataset.morph || null, tag: el.tagName.toLowerCase() + (el.className ? '.' + String(el.className).split(' ')[0] : '') }; }
  function boxItems(el) {
    const cs = getComputedStyle(el); if (el.tagName === 'svg' || el.closest('svg')) return;
    const b = box(el); if (b.w < 0.5 || b.h < 0.5) return;
    const clipText = (cs.backgroundClip || cs.webkitBackgroundClip) === 'text';     // gradient *text* (.grad): no box, the runs carry the colour
    const fill = clipText ? null : rgba(cs.backgroundColor); const grad = !clipText && cs.backgroundImage !== 'none';
    const bw = [cs.borderLeftWidth, cs.borderTopWidth, cs.borderRightWidth, cs.borderBottomWidth].map(num);
    const bc = [cs.borderLeftColor, cs.borderTopColor, cs.borderRightColor, cs.borderBottomColor].map(rgba);
    const bs = [cs.borderLeftStyle, cs.borderTopStyle, cs.borderRightStyle, cs.borderBottomStyle];
    const side = i => bw[i] > 0 && bc[i] && bs[i] !== 'none' && bs[i] !== 'hidden';
    const uniform = side(0) && side(1) && side(2) && side(3) && bw.every(v => v === bw[0]) && bc.every(c => c.hex === bc[0].hex);
    const rad = cs.borderTopLeftRadius.endsWith('%') ? Math.min(b.w, b.h) * num(cs.borderTopLeftRadius) / 100 : num(cs.borderTopLeftRadius);
    if (fill || grad || uniform) items.push({ type: 'box', id: flag(el), x: b.x, y: b.y, w: b.w, h: b.h, fill, grad, radius: rad,
        line: uniform ? { color: bc[0], w: bw[0], dash: bs[0] === 'dashed' } : null, morph: el.dataset.morph || null, tag: el.tagName.toLowerCase() + '.' + String(el.className || '').split(' ')[0] });
    else if (side(0) || side(1) || side(2) || side(3)) flag(el);
    if (!uniform) { // per-side borders as thin rectangles
      if (side(1)) items.push({ type: 'box', x: b.x, y: b.y, w: b.w, h: bw[1], fill: bc[1], radius: 0, line: null });
      if (side(3)) items.push({ type: 'box', x: b.x, y: b.y + b.h - bw[3], w: b.w, h: bw[3], fill: bc[3], radius: 0, line: null });
      if (side(0)) items.push({ type: 'box', x: b.x, y: b.y, w: bw[0], h: b.h, fill: bc[0], radius: 0, line: null });
      if (side(2)) items.push({ type: 'box', x: b.x + b.w - bw[2], y: b.y, w: bw[2], h: b.h, fill: bc[2], radius: 0, line: null }); }
    // ::before decorations (the bullet dashes of ul.big li)
    const ps = getComputedStyle(el, '::before');
    if (ps.content !== 'none' && ps.content !== 'normal' && ps.position === 'absolute' && (rgba(ps.backgroundColor) || ps.backgroundImage !== 'none') && num(ps.width) > 0) {
      const fs = num(cs.fontSize), top = ps.top.endsWith('em') ? num(ps.top) * fs : num(ps.top);
      items.push({ type: 'box', x: b.x + num(ps.left), y: b.y + top, w: num(ps.width), h: num(ps.height), fill: rgba(ps.backgroundColor), grad: ps.backgroundImage !== 'none', radius: num(ps.borderTopLeftRadius), line: null }); } }
  function figure(svg) {
    const b = box(svg); const id = flag(svg);
    const hide = [];
    // 1) text labels -> native text boxes (skip rotated ones; they stay in the picture)
    const texts = [];
    svg.querySelectorAll('text').forEach(t => {
      if (!visible(t)) return; if ((t.getAttribute('transform') || '').includes('rotate')) return;
      const r = t.getBoundingClientRect(); if (r.width < 0.5) return; const cs = getComputedStyle(t);
      const st = style(cs, null); st.color = rgba(cs.fill) || st.color;
      texts.push({ type: 'text', id: flag(t), x: r.left, y: r.top, w: r.width, h: r.height, align: 'left', lh: r.height, pre: false, wrap: false,
                   runs: [Object.assign({ t: t.textContent }, st)], morph: t.dataset.morph || null, tag: 'svg-text', svgtext: true }); hide.push(t); });
    // 2) geometry that morphs to a neighbouring slide -> its own cut-out picture
    const cuts = [];
    svg.querySelectorAll('[data-morph]').forEach(e => {
      if (e.tagName === 'text' || !want.has(e.dataset.morph) || !visible(e)) return;
      const r = e.getBoundingClientRect(); if (r.width < 0.5 || r.height < 0.5) return;
      cuts.push({ type: 'image', kind: 'cut', id: flag(e), x: r.left, y: r.top, w: r.width, h: r.height, pad: 6, morph: e.dataset.morph }); hide.push(e); });
    items.push({ type: 'image', kind: 'fig', id, x: b.x, y: b.y, w: b.w, h: b.h, pad: 4, hide: hide.map(e => e.getAttribute('data-x')), morph: null });
    items.push(...cuts, ...texts); }
  function walk(el) {
    if (el.tagName === 'ASIDE' || !visible(el)) return;
    if (el.classList.contains('glow')) {            // soft radial glow -> ellipse with a radial gradient fading to transparent
      const b = box(el), cs = getComputedStyle(el), c = rgba((cs.backgroundImage.match(/rgba?\([^)]*\)/) || [''])[0]);
      if (c) items.push({ type: 'glow', id: flag(el), x: b.x - b.w * 0.1, y: b.y - b.h * 0.1, w: b.w * 1.2, h: b.h * 1.2, color: c }); else flag(el);
      return; }
    if (el.tagName === 'TABLE') flag(el);           // cells and rules are emitted individually; keep collapsed borders out of the residual
    if (el.tagName === 'svg') { if (el.classList.contains('fig')) figure(el); else { const b = box(el); items.push({ type: 'image', kind: 'logo', id: flag(el), x: b.x, y: b.y, w: b.w, h: b.h }); } return; }
    boxItems(el);
    if (isTextContainer(el)) { items.push(textItem(el)); return; }
    for (const n of el.childNodes) {
      if (n.nodeType === 1) walk(n);
      else if (n.nodeType === 3 && n.textContent.trim()) {   // a bare text node next to block children (chip labels)
        const rg = document.createRange(); rg.selectNodeContents(n); const r = rg.getBoundingClientRect();
        const cs = getComputedStyle(el); const st = style(cs, null);
        items.push({ type: 'text', x: r.left, y: r.top, w: r.width, h: r.height, align: 'left', lh: r.height, pre: false, wrap: false,
                     runs: [Object.assign({ t: n.textContent.replace(/\s+/g, ' ').trim() }, st)], morph: null, tag: 'textnode', hidewith: flag(el) }); } } }
  slide.querySelectorAll('[data-x]').forEach(e => e.removeAttribute('data-x'));
  for (const n of slide.children) if (n.nodeType === 1) walk(n);
  const root = getComputedStyle(document.documentElement);
  return { items, acc: root.getPropertyValue('--acc1').trim(), accDeep: root.getPropertyValue('--acc-deep').trim(), bg: root.getPropertyValue('--bg').trim() };
}"""

# ---- helpers for the PPTX side --------------------------------------------------------------------
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006", "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
      "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main"}
E = lambda px: Emu(int(round(px * PX)))
hexcol = lambda s: s.lstrip("#").upper()

def set_transition(slide, morph):
    sld = slide._element
    for old in sld.findall(qn("p:transition")) + sld.findall(f"{{{NS['mc']}}}AlternateContent"): sld.remove(old)
    def fade():
        f = etree.Element(qn("p:transition"), nsmap={"p14": NS["p14"]}); f.set("spd", "slow"); f.set(f"{{{NS['p14']}}}dur", str(DUR_MS))
        etree.SubElement(f, qn("p:fade")); return f
    if morph:
        ac = etree.Element(f"{{{NS['mc']}}}AlternateContent", nsmap={"mc": NS["mc"]})
        ch = etree.SubElement(ac, f"{{{NS['mc']}}}Choice", nsmap={"p159": NS["p159"]}); ch.set("Requires", "p159")
        tr = etree.SubElement(ch, qn("p:transition"), nsmap={"p14": NS["p14"]}); tr.set("spd", "slow"); tr.set(f"{{{NS['p14']}}}dur", str(DUR_MS))
        etree.SubElement(tr, f"{{{NS['p159']}}}morph").set("option", "byObject")
        etree.SubElement(ac, f"{{{NS['mc']}}}Fallback").append(fade()); node = ac
    else: node = fade()
    anchor = sld.find(qn("p:clrMapOvr")); (anchor if anchor is not None else sld.find(qn("p:cSld"))).addnext(node)

def alpha(clr_el, a):
    if a < 0.999: etree.SubElement(clr_el, qn("a:alpha")).set("val", str(int(round(a * 100000))))

def add_box(slide, it, acc, name):
    r = it.get("radius", 0) or 0; w, h = max(it["w"], 0.5), max(it["h"], 0.5)
    circle = r > 0 and r >= min(w, h) / 2 - 0.5 and abs(w - h) < 1.5          # a pill (border-radius 999px) is a rounded rectangle, not an ellipse
    kind = MSO_SHAPE.OVAL if circle else (MSO_SHAPE.ROUNDED_RECTANGLE if r > 0 else MSO_SHAPE.RECTANGLE)
    sh = slide.shapes.add_shape(kind, E(it["x"]), E(it["y"]), E(w), E(h)); sh.name = name
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = min(0.5, r / min(w, h))
    sh.shadow.inherit = False
    fill = it.get("fill")
    if it.get("grad") and not fill: fill = {"hex": acc, "a": 1}
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill["hex"])
        alpha(sh._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), fill["a"])
    else: sh.fill.background()
    ln = it.get("line")
    if ln and ln.get("color"):
        sh.line.color.rgb = RGBColor.from_string(ln["color"]["hex"]); sh.line.width = Pt(max(ln["w"], 1) * PT)
        alpha(sh.line._get_or_add_ln().find(qn("a:solidFill")).find(qn("a:srgbClr")), ln["color"]["a"])
        if ln.get("dash"): sh.line.dash_style = MSO_LINE.DASH
    else: sh.line.fill.background()
    # shapes must not carry text-frame defaults that could show up
    tf = sh.text_frame; tf.text = ""
    return sh

def add_glow(slide, it, name):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"])); sh.name = name
    sh.shadow.inherit = False; sh.line.fill.background(); sh.text_frame.text = ""
    spPr = sh._element.spPr
    for old in spPr.findall(qn("a:solidFill")) + spPr.findall(qn("a:gradFill")) + spPr.findall(qn("a:noFill")): spPr.remove(old)
    gf = etree.Element(qn("a:gradFill")); gf.set("rotWithShape", "1"); gs = etree.SubElement(gf, qn("a:gsLst"))
    for pos, a in ((0, it["color"]["a"]), (70000, 0.0), (100000, 0.0)):
        g = etree.SubElement(gs, qn("a:gs")); g.set("pos", str(pos)); c = etree.SubElement(g, qn("a:srgbClr")); c.set("val", it["color"]["hex"])
        etree.SubElement(c, qn("a:alpha")).set("val", str(int(round(a * 100000))))
    path = etree.SubElement(gf, qn("a:path")); path.set("path", "circle")
    ftr = etree.SubElement(path, qn("a:fillToRect")); [ftr.set(k, "50000") for k in ("l", "t", "r", "b")]
    spPr.find(qn("a:prstGeom")).addnext(gf)
    return sh

def add_text(slide, it, acc, acc_deep, name):
    # a text that fits on one line in the browser must never wrap in PowerPoint (its metrics differ by a hair);
    # multi-line text keeps wrapping, with a little slack so the break points stay where the browser put them
    single = it["h"] <= it["lh"] * 1.6 or not it.get("wrap")
    w = max(it["w"], 1) + (3 if single else 2)
    x = it["x"] - (1 if it.get("align") == "center" else 0)
    tb = slide.shapes.add_textbox(E(x), E(it["y"]), E(w), E(max(it["h"], 1))); tb.name = name
    tf = tb.text_frame; tf.word_wrap = not single; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.TOP
    # split runs into paragraphs on <br> and, for <pre>, on newlines
    paras, cur = [], []
    for r in it["runs"]:
        if r.get("br"): paras.append(cur); cur = []; continue
        parts = r["t"].split("\n") if it.get("pre") else [r["t"]]
        for j, part in enumerate(parts):
            if j: paras.append(cur); cur = []
            if part: cur.append(dict(r, t=part))
    paras.append(cur)
    # trim whitespace at paragraph edges (not in <pre>, where indentation is content), drop empty paragraphs at the ends
    for p in paras:
        if p and not it.get("pre"): p[0]["t"] = p[0]["t"].lstrip(); p[-1]["t"] = p[-1]["t"].rstrip()
    while paras and not any(r["t"] for r in paras[0]): paras.pop(0)
    while paras and not any(r["t"] for r in paras[-1]): paras.pop()
    if not paras: paras = [[]]
    align = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[it.get("align", "left")]
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(max(it["lh"], 1) * PT) if it.get("svgtext") is not True else 1.0
        for r in runs:
            if not r["t"]: continue
            run = p.add_run(); run.text = r["t"]; f = run.font
            f.size = Pt(max(r["size"], 4) * PT)      # the computed size already includes <sup>'s smaller font-size
            f.bold = bool(r.get("bold")); f.italic = bool(r.get("italic")); f.name = MONO if r.get("mono") else SANS
            rPr = run._r.get_or_add_rPr()
            if r.get("grad"):                      # gradient text (.grad) -> solid scarlet: gradient text fill is PowerPoint-only
                f.color.rgb = RGBColor.from_string(acc)
            elif r.get("color"):
                f.color.rgb = RGBColor.from_string(r["color"]["hex"]); alpha(rPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), r["color"]["a"])
            if r.get("ls"): rPr.set("spc", str(int(round(r["ls"] * PT * 100))))
            if r.get("caps"): rPr.set("cap", "all")
            if r.get("sup"): rPr.set("baseline", "30000")
            if r.get("sub"): rPr.set("baseline", "-25000")
            # rPr attribute order does not matter, but the children do: solidFill/gradFill before latin
            latin = rPr.find(qn("a:latin"))
            if latin is not None: rPr.remove(latin); rPr.append(latin)
    return tb

# ---- render / extract with headless Chromium -----------------------------------------------------
tmp = pathlib.Path(tempfile.mkdtemp(prefix="deck-pptx-"))
slides_data = []
with sync_playwright() as p:
    browser = p.chromium.launch()
    ctx = browser.new_context(viewport={"width": W, "height": H}, device_scale_factor=1, color_scheme=theme, reduced_motion="reduce")
    ctx.add_init_script(f"try{{localStorage.setItem('deck-theme','{theme}')}}catch(e){{}}")
    page = ctx.new_page(); page.goto(src.as_uri() + "#/1")
    page.add_style_tag(content="#progress,#hint,#counter,#help,#overview,#notes{display:none!important}")
    page.evaluate("document.head.appendChild(Object.assign(document.createElement('style'),{id:'xstyle'}))")
    page.wait_for_timeout(600)
    set_style = lambda css: page.evaluate("c => { document.getElementById('xstyle').textContent = c }", css)
    TRANSPARENT = "html,body,#viewport,#stage,.slide{background:transparent!important}.glow{display:none!important}"
    def shoot(path, x, y, w, h, pad=0, transparent=True):
        x0, y0 = max(0, int(x) - pad), max(0, int(y) - pad); x1, y1 = min(W, int(x + w) + pad + 1), min(H, int(y + h) + pad + 1)
        if x1 - x0 < 1 or y1 - y0 < 1: return None
        page.screenshot(path=str(path), omit_background=transparent, clip={"x": x0, "y": y0, "width": x1 - x0, "height": y1 - y0})
        return (x0, y0, x1 - x0, y1 - y0)
    ids = [set(x) for x in page.evaluate("() => Array.from(document.querySelectorAll('#stage > .slide')).map(s => Array.from(s.querySelectorAll('[data-morph]')).map(e => e.dataset.morph))")]

    # a crisp Double T once (3x) for every logo placement
    logo_png = tmp / "ttu-dt.png"
    hi = browser.new_context(viewport={"width": 400, "height": 400}, device_scale_factor=3, color_scheme="light")
    hp = hi.new_page(); hp.goto(src.as_uri() + "#/1"); hp.wait_for_timeout(500)
    hp.evaluate("""() => { const sym = document.getElementById('ttu-dt'); document.body.innerHTML = ''; document.body.style.background = 'transparent'; document.documentElement.style.background = 'transparent';
        const s = document.createElementNS('http://www.w3.org/2000/svg','svg'); s.setAttribute('viewBox','0 0 85.9 100.7'); s.style.cssText='width:344px;height:403px;display:block;position:absolute;left:0;top:0';
        s.appendChild(sym.cloneNode(true)); const u = document.createElementNS('http://www.w3.org/2000/svg','use'); u.setAttribute('href','#ttu-dt'); s.appendChild(u); document.body.appendChild(s); }""")
    hp.wait_for_timeout(200); hp.screenshot(path=str(logo_png), omit_background=True, clip={"x": 0, "y": 0, "width": 344, "height": 403}); hi.close()

    for i in range(N):
        page.evaluate(f"window.deck.go({i})"); page.wait_for_timeout(600); page.wait_for_function("!document.querySelector('.slide.leaving')")
        want = set()
        if MORPH:
            if i > 0 and share[i]: want |= ids[i] & ids[i - 1]
            if i + 1 < N and share[i + 1]: want |= ids[i] & ids[i + 1]
        set_style("")
        if mode == "image":
            boxes = page.evaluate("""want => { const s = document.querySelector('.slide.active'), out = [], seen = new Set();
                s.querySelectorAll('[data-x]').forEach(e => e.removeAttribute('data-x'));
                s.querySelectorAll('[data-morph]').forEach(e => { const id = e.dataset.morph; if (!want.includes(id) || seen.has(id)) return; seen.add(id);
                  const r = e.getBoundingClientRect(); if (r.width < 1 || r.height < 1) return; e.setAttribute('data-x', id); out.push({id, x: r.left, y: r.top, w: r.width, h: r.height}); });
                return out; }""", sorted(want))
            set_style(".slide.active [data-x]{visibility:hidden!important}")
            bg = tmp / f"bg-{i + 1:02d}.png"; page.screenshot(path=str(bg), clip={"x": 0, "y": 0, "width": W, "height": H})
            parts = []
            for b in boxes:
                sel = f'.slide.active [data-x="{b["id"]}"]'
                set_style(TRANSPARENT + f".slide.active *{{visibility:hidden!important}}{sel},{sel} *{{visibility:visible!important}}")
                f = tmp / f"el-{i + 1:02d}-{len(parts):03d}.png"; g = shoot(f, b["x"], b["y"], b["w"], b["h"], 6)
                if g: parts.append({"type": "image", "kind": "cut", "morph": b["id"], "file": f, "geom": g})
            slides_data.append({"bg": bg, "items": parts}); print(f"slide {i + 1:2d}: {len(parts):2d} morph parts"); continue

        data = page.evaluate(EXTRACT_JS, {"want": sorted(want)})
        items = data["items"]
        # pictures: figures (with lifted labels / morph pieces hidden), cut-outs, logos
        for it in items:
            if it["type"] != "image": continue
            sel = f'.slide.active [data-x="{it["id"]}"]'
            hide = "".join(f'.slide.active [data-x="{h}"],.slide.active [data-x="{h}"] *{{visibility:hidden!important}}' for h in it.get("hide", []))
            if it["kind"] == "logo":
                it["file"] = logo_png; it["geom"] = (it["x"], it["y"], it["w"], it["h"]); continue
            set_style(TRANSPARENT + f".slide.active *{{visibility:hidden!important}}{sel},{sel} *{{visibility:visible!important}}" + hide)
            f = tmp / f"img-{i + 1:02d}-{it['id']}.png"; it["geom"] = shoot(f, it["x"], it["y"], it["w"], it["h"], it.get("pad", 0)); it["file"] = f
        # residual layer: whatever the walk did not turn into an object (the soft glows, mostly)
        set_style("html,body,#viewport,#stage,.slide{background:transparent!important}.slide.active [data-x]{visibility:hidden!important}")
        res = tmp / f"res-{i + 1:02d}.png"; page.screenshot(path=str(res), omit_background=True, clip={"x": 0, "y": 0, "width": W, "height": H})
        set_style("")
        slides_data.append({"items": items, "residual": res, "acc": hexcol(data["acc"]), "accDeep": hexcol(data["accDeep"]), "bg": hexcol(data["bg"])})
        nt = sum(1 for it in items if it["type"] == "text"); nb = sum(1 for it in items if it["type"] == "box"); ni = sum(1 for it in items if it["type"] == "image")
        print(f"slide {i + 1:2d}: {nt:3d} text boxes, {nb:3d} shapes, {ni:2d} pictures" + (f", morph ids: {len(want)}" if want else ""))
    browser.close()

# ---- assemble --------------------------------------------------------------------------------------
from PIL import Image
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
def set_alt(shape, text): shape._element.xpath("./*[local-name()='nvSpPr' or local-name()='nvPicPr']/*[local-name()='cNvPr']")[0].set("descr", text)
def residual_nonempty(path):
    im = Image.open(path).convert("RGBA"); a = im.getchannel("A"); return a.getbbox() is not None and sum(a.histogram()[8:]) > 200

for i, (m, d) in enumerate(zip(meta, slides_data)):
    slide = prs.slides.add_slide(blank)
    if mode == "image":
        pic = slide.shapes.add_picture(str(d["bg"]), 0, 0, prs.slide_width, prs.slide_height); pic.name = f"Slide {i + 1} background"; set_alt(pic, f"Slide {i + 1} of {N}: {m['title']}")
        for it in d["items"]:
            x, y, w, h = it["geom"]; el = slide.shapes.add_picture(str(it["file"]), E(x), E(y), E(w), E(h)); el.name = f"!!{it['morph']}"
    else:
        # solid slide background
        bgf = slide.background.fill; bgf.solid(); bgf.fore_color.rgb = RGBColor.from_string(d["bg"])
        if residual_nonempty(d["residual"]):
            im = Image.open(d["residual"]).convert("RGBA"); a = im.getchannel("A"); strong = a.point(lambda v: 255 if v >= 128 else 0).getbbox()
            pic = slide.shapes.add_picture(str(d["residual"]), 0, 0, prs.slide_width, prs.slide_height); pic.name = f"s{i + 1}-backdrop"; set_alt(pic, "decorative backdrop")
            if strong: print(f"  slide {i + 1}: backdrop keeps opaque pixels in {strong} (something was not converted)")
        counts = {"box": 0, "text": 0, "image": 0, "glow": 0}
        for it in d["items"]:
            counts[it["type"]] += 1
            name = f"!!{it['morph']}" if it.get("morph") else f"s{i + 1}-{it['type']}-{counts[it['type']]:02d}"
            if it["type"] == "glow": add_glow(slide, it, name)
            elif it["type"] == "box": add_box(slide, it, d["acc"], name)
            elif it["type"] == "text": add_text(slide, it, d["acc"], d["accDeep"], name)
            else:
                if not it.get("geom"): continue
                x, y, w, h = it["geom"]; pic = slide.shapes.add_picture(str(it["file"]), E(x), E(y), E(w), E(h)); pic.name = name
                set_alt(pic, {"fig": "diagram", "cut": "chart element", "logo": "Texas Tech University Double T"}[it["kind"]])
    if m["notes"]: slide.notes_slide.notes_text_frame.text = m["notes"]
    if i > 0: set_transition(slide, MORPH and share[i])
prs.core_properties.title = doc_title; prs.core_properties.author = "Ziqing Guo"
prs.core_properties.subject = "Ph.D. defense, Texas Tech University, October 9, 2026"
out.parent.mkdir(parents=True, exist_ok=True); prs.save(str(out))
n_morph = sum(1 for i in range(1, N) if MORPH and share[i])
print(f"wrote {out} ({out.stat().st_size // 1024} KB): {N} slides, mode={mode}, {n_morph} Morph + {N - 1 - n_morph} Fade transitions, "
      f"{sum(1 for m in meta if m['notes'])} with notes, theme={theme}")
